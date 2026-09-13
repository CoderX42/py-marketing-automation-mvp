from __future__ import annotations

from dataclasses import dataclass
import json
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class InputRecord:
    row_number: int
    phone: str
    name: str = ""


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="1F2937")
SUCCESS_FILL = PatternFill("solid", fgColor="E2F0D9")
FAILED_FILL = PatternFill("solid", fgColor="FCE4D6")
MAX_CELL_TEXT = 30000
STATUS_LABELS = {
    "success": "已完成", "empty": "无结果", "failed": "失败",
    "partial": "采集不完整", "login_required": "等待重新登录", "pending": "待处理",
}


def _append_values(ws, values) -> None:
    """Always store source strings literally, never as Excel formulas."""
    values = list(values)
    for value in values:
        if isinstance(value, str) and len(value.encode("utf-16-le")) // 2 > 32767:
            raise ValueError("表格字段超过 Excel 单元格容量，尚未保存；请保留结果备份")
    ws.append(values)
    # Avoid ws.max_row / ws[row] here: those scan the whole sheet on every append.
    row_number = ws._current_row
    for column, value in enumerate(values, start=1):
        if isinstance(value, str):
            ws.cell(row_number, column).data_type = "s"


def _text_chunks(text: str) -> list[str]:
    # Count UTF-16 units as well as characters so emoji cannot exceed Excel's limit.
    chunks, start, units = [], 0, 0
    for index, char in enumerate(text):
        char_units = 2 if ord(char) > 0xFFFF else 1
        if units + char_units > MAX_CELL_TEXT:
            chunks.append(text[start:index])
            start, units = index, 0
        units += char_units
    chunks.append(text[start:])
    return chunks


def _style_sheet(ws, widths: dict[str, float], wrap_columns: set[str] | None = None) -> None:
    wrap_columns = wrap_columns or set()
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=get_column_letter(cell.column) in wrap_columns)
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions


def create_input_template(path: str | Path) -> Path:
    """Create the user-facing phone list template used by the importer."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "手机号清单"
    ws.append(["手机号", "姓名"])
    ws.append(["", ""])
    _style_sheet(ws, {"A": 20, "B": 18})
    for cell in ws["A"]:
        cell.number_format = "@"
    notes = wb.create_sheet("使用说明")
    notes.append(["填写说明"])
    notes.append(["1. 在“手机号清单”中从第 2 行开始填写手机号。"])
    notes.append(["2. 手机号请保持 11 位数字，姓名可以留空。"])
    notes.append(["3. 保存后在营销详情批量查询工具中选择此文件。"])
    _style_sheet(notes, {"A": 72}, {"A"})
    target_path = target.with_suffix(".xlsx")
    wb.save(target_path)
    return target_path


def normalize_phone(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text.replace(" ", "").replace("-", "")


def read_input(path: str | Path, phone_column: int = 1, name_column: int = 2) -> list[InputRecord]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []

    # 工作簿200没有表头；带“手机号/手机号码”的文件自动跳过首行。
    start = 1
    first = [str(v or "").strip() for v in rows[0]]
    if any(v in {"手机号", "手机号码", "电话", "电话号码"} for v in first):
        start = 2

    records: list[InputRecord] = []
    seen: set[str] = set()
    for number, row in enumerate(rows[start - 1 :], start=start):
        phone = normalize_phone(row[phone_column - 1] if len(row) >= phone_column else "")
        name = str(row[name_column - 1] or "").strip() if len(row) >= name_column else ""
        if not phone or phone in seen:
            continue
        if len(phone) != 11 or not phone.isdigit() or not phone.startswith("1"):
            continue
        seen.add(phone)
        records.append(InputRecord(number, phone, name))
    return records


def create_result_workbook(records: Iterable[InputRecord]) -> Workbook:
    wb = Workbook()
    source = wb.active
    source.title = "输入数据"
    source.append(["手机号", "姓名"])
    for record in records:
        _append_values(source, [record.phone, record.name])
    for cell in source["A"]:
        cell.number_format = "@"
    _style_sheet(source, {"A": 18, "B": 16})

    summary = wb.create_sheet("结果汇总")
    summary.append(["手机号", "姓名", "处理状态", "查询时间", "推荐数量", "推荐摘要", "错误信息"])
    detail_text = wb.create_sheet("营销详情")
    detail_text.append(["手机号", "姓名", "处理状态", "查询时间", "详情分段", "营销详情全文", "错误信息"])
    detail = wb.create_sheet("推荐明细")
    detail.append(["手机号", "推荐序号", "业务名称", "资费或承诺", "产品描述", "原始文本", "推荐分段"])
    log = wb.create_sheet("运行日志")
    log.append(["时间", "手机号", "事件", "详情"])
    _style_sheet(summary, {"A": 18, "B": 16, "C": 12, "D": 22, "E": 12, "F": 42, "G": 48}, {"F", "G"})
    _style_sheet(detail_text, {"A": 18, "B": 16, "C": 12, "D": 22, "E": 10, "F": 100, "G": 48}, {"F", "G"})
    _style_sheet(detail, {"A": 18, "B": 12, "C": 42, "D": 24, "E": 60, "F": 100, "G": 10}, {"C", "D", "E", "F"})
    _style_sheet(log, {"A": 22, "B": 18, "C": 14, "D": 72}, {"D"})
    # Open the workbook on the useful reader-facing summary first.
    wb._sheets = [summary, detail_text, detail, source, log]
    return wb


def write_result(path: str | Path, records: list[InputRecord], results: dict[str, dict]) -> Path:
    wb = create_result_workbook(records)
    source = wb["输入数据"]
    summary = wb["结果汇总"]
    detail_text = wb["营销详情"]
    detail = wb["推荐明细"]
    log = wb["运行日志"]
    for record in records:
        result = results.get(record.phone, {})
        # Older saved results may still use the ISO date/time separator.
        query_time = str(result.get("query_time") or "").replace("T", " ", 1)
        items = result.get("items", [])
        raw_text = result.get("raw_text", "") or ""
        status = result.get("status", "pending")
        display_status = STATUS_LABELS.get(status, status)
        summary_text = "；".join(item.get("name", "") for item in items)
        if not summary_text and raw_text:
            summary_text = raw_text.replace("\n", " ")[:300]
        _append_values(summary, [
            record.phone,
            record.name,
            display_status,
            query_time,
            len(items),
            summary_text[:1000],
            result.get("error", ""),
        ])
        # Excel 单元格最多保存 32,767 个字符；用 30,000 字符分段，确保
        # 超长营销详情不会被 openpyxl/Excel 静默截断。每段保留号码和段号，
        # 使用者可在“营销详情”页按号码和段号筛选、排序后阅读完整内容。
        chunks = _text_chunks(raw_text)
        for part, chunk in enumerate(chunks, start=1):
            _append_values(detail_text, [
                record.phone,
                record.name,
                display_status,
                query_time,
                part,
                chunk,
                result.get("error", "") if part == 1 else "",
            ])
        for index, item in enumerate(items, start=1):
            fields = [_text_chunks(str(item.get(key, "") or "")) for key in ("name", "price", "description", "raw_text")]
            for part in range(max(map(len, fields))):
                _append_values(detail, [record.phone, index, *[field[part] if part < len(field) else "" for field in fields], part + 1])
        _append_values(log, [query_time or datetime.now().isoformat(sep=" ", timespec="seconds"), record.phone, display_status, result.get("error", "")])
    for ws in (source, summary, detail_text, detail, log):
        wrap = {"输入数据": {"B"}, "结果汇总": {"F", "G"}, "营销详情": {"F", "G"}, "推荐明细": {"C", "D", "E", "F"}, "运行日志": {"D"}}[ws.title]
        _style_sheet(ws, {}, wrap)
        phone_column = "B" if ws.title == "运行日志" else "A"
        for cell in ws[phone_column]:
            cell.number_format = "@"
        for row in ws.iter_rows(min_row=2):
            if ws.title in {"结果汇总", "营销详情", "运行日志"}:
                cell = row[2]
                if cell.value != STATUS_LABELS["pending"]:
                    cell.fill = SUCCESS_FILL if cell.value in {STATUS_LABELS["success"], STATUS_LABELS["empty"]} else FAILED_FILL
            if ws.title in {"营销详情", "推荐明细"}:
                text = str(row[5].value or "")
                # Excel limits row heights; leave wrapped text readable and the
                # complete cell text available in the formula bar / editing view.
                lines = max(text.count("\n") + 1, len(text) // 65 + 1)
                ws.row_dimensions[row[0].row].height = min(300, max(48, lines * 16))
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    # Replace the previous workbook atomically so an interruption during save
    # cannot leave a half-written result file.
    fd, temporary_name = tempfile.mkstemp(prefix=f".{target.stem}-", suffix=".xlsx", dir=target.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        wb.save(temporary)
        with temporary.open("r+b") as saved:
            os.fsync(saved.fileno())
        os.replace(temporary, target)
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)
    return target


def load_existing_results(path: str | Path) -> dict[str, dict]:
    """Restore all statuses and complete text; the caller decides which to retry."""
    file_path = Path(path)
    if not file_path.exists():
        return {}
    wb = load_workbook(file_path, read_only=True, data_only=False)
    output: dict[str, dict] = {}
    status_keys = {label: key for key, label in STATUS_LABELS.items()}

    def rows(sheet):
        if sheet not in wb.sheetnames:
            return
        iterator = wb[sheet].iter_rows(values_only=True)
        headers = next(iterator, ())
        for row in iterator:
            yield {header: value for header, value in zip(headers, row) if header}

    try:
        for row in rows("结果汇总"):
            phone = normalize_phone(row.get("手机号"))
            if not phone:
                continue
            status = str(row.get("处理状态") or "pending")
            output[phone] = {
                "status": status_keys.get(status, status), "items": [], "raw_text": "",
                "query_time": str(row.get("查询时间") or ""), "error": str(row.get("错误信息") or ""),
            }
        parts: dict[str, list[tuple[int, str]]] = {}
        for row in rows("营销详情"):
            phone = normalize_phone(row.get("手机号"))
            if phone not in output:
                continue
            # Older workbooks have one unnumbered full-text row per phone.
            parts.setdefault(phone, []).append((int(row.get("详情分段") or 1), str(row.get("营销详情全文") or "")))
        for phone, segments in parts.items():
            output[phone]["raw_text"] = "".join(text for _, text in sorted(segments, key=lambda entry: entry[0]))

        item_parts: dict[tuple[str, int], list[tuple[int, dict]]] = {}
        for row in rows("推荐明细"):
            phone = normalize_phone(row.get("手机号"))
            if phone not in output:
                # A detail row is not proof that a phone was successfully queried.
                continue
            index = int(row.get("推荐序号") or 1)
            item = {key: str(row.get(header) or "") for key, header in (
                ("name", "业务名称"), ("price", "资费或承诺"), ("description", "产品描述"), ("raw_text", "原始文本"))}
            item_parts.setdefault((phone, index), []).append((int(row.get("推荐分段") or 1), item))
        for (phone, index), segments in sorted(item_parts.items()):
            segments.sort(key=lambda entry: entry[0])
            output[phone]["items"].append({key: "".join(item[key] for _, item in segments) for key in ("name", "price", "description", "raw_text")})
    finally:
        wb.close()
    return output


def _journal_path(path: str | Path) -> Path:
    return Path(path).with_suffix(".jsonl")


def append_result_journal(path: str | Path, phone: str, result: dict) -> Path:
    """Durably save a completed/partial result before attempting Excel replacement.

    The journal sits next to the workbook with the same stem and .jsonl suffix.
    Errors propagate: a caller must not continue querying if this write fails.
    """
    phone = normalize_phone(phone)
    if not phone or not isinstance(result, dict):
        raise ValueError("备份结果必须包含手机号和结果对象")
    entry = {"schema_version": 1, "phone": phone, "result": result}
    payload = (json.dumps(entry, ensure_ascii=False, allow_nan=False) + "\n").encode("utf-8")
    target = _journal_path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    fd = os.open(target, os.O_RDWR | os.O_CREAT | os.O_APPEND, 0o600)
    with os.fdopen(fd, "a+b") as stream:
        # An interrupted append may leave an incomplete final line. Remove only
        # that unfinished line, retaining all prior fsynced query results.
        stream.seek(0, os.SEEK_END)
        end = stream.tell()
        if end:
            stream.seek(end - 1)
            if stream.read(1) != b"\n":
                position, suffix = end, b""
                while position:
                    length = min(position, 8192)
                    position -= length
                    stream.seek(position)
                    suffix = stream.read(length) + suffix
                    if b"\n" in suffix:
                        break
                tail_start = position + suffix.rfind(b"\n") + 1 if b"\n" in suffix else 0
                tail = suffix.rsplit(b"\n", 1)[-1]
                try:
                    json.loads(tail)
                except (ValueError, UnicodeDecodeError):
                    stream.truncate(tail_start)
                else:
                    stream.write(b"\n")
        stream.write(payload)
        stream.flush()
        os.fsync(stream.fileno())
    return target


def load_result_journal(path: str | Path) -> dict[str, dict]:
    """Recover the latest result per phone; ignore only a truncated final line."""
    target = _journal_path(path)
    if not target.exists():
        return {}
    output = {}
    with target.open("rb") as stream:
        for number, line in enumerate(stream, start=1):
            if not line.strip():
                continue
            try:
                entry = json.loads(line)
            except (ValueError, UnicodeDecodeError) as exc:
                if not line.endswith(b"\n"):
                    break
                raise ValueError(f"结果备份第 {number} 行损坏，已保留原文件") from exc
            if not isinstance(entry, dict) or entry.get("schema_version") != 1 or not isinstance(entry.get("result"), dict):
                raise ValueError(f"结果备份第 {number} 行格式不受支持，已保留原文件")
            phone = normalize_phone(entry.get("phone"))
            if not phone:
                raise ValueError(f"结果备份第 {number} 行缺少手机号，已保留原文件")
            output[phone] = entry["result"]
    return output
