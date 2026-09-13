"""Export durability regression tests using synthetic numbers and text only."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from excel_io import (
    InputRecord, append_result_journal, load_existing_results,
    load_result_journal, write_result,
)


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "查询结果.xlsx"
        self.phone = "13800000000"
        self.records = [InputRecord(2, self.phone, "测试姓名")]

    def result(self, status="success", text="营销详情\n重复标题\n重复标题\n属地推荐方案"):
        return {
            "status": status, "query_time": "2026-09-13 12:00:00", "raw_text": text,
            "error": "" if status == "success" else "测试异常",
            "items": [{"name": "测试推荐", "price": "每月 10 元", "description": "测试描述", "raw_text": text}],
        }

    def test_multichunk_fulltext_and_items_survive_repeated_resume(self):
        text = ("重复标题\n方案 😀\n" * 8000) + "属地推荐方案\n最后一行"
        original = self.result(text=text)
        self.assertEqual(write_result(self.path, self.records, {self.phone: original}), self.path)
        restored = load_existing_results(self.path)
        self.assertEqual(restored[self.phone], original)
        write_result(self.path, self.records, restored)
        self.assertEqual(load_existing_results(self.path)[self.phone], original)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertGreater(wb["营销详情"].max_row, 2)
        for row in wb["营销详情"].iter_rows(min_row=2):
            self.assertLessEqual(len(row[5].value.encode("utf-16-le")) // 2, 30000)
            self.assertTrue(row[5].alignment.wrap_text)
        self.assertEqual(wb["营销详情"].auto_filter.ref, wb["营销详情"].dimensions)
        self.assertEqual(wb.active.title, "结果汇总")

    def test_failure_with_detail_items_is_not_promoted_to_success(self):
        for status in ("partial", "failed", "login_required", "empty", "pending"):
            with self.subTest(status=status):
                original = self.result(status=status)
                write_result(self.path, self.records, {self.phone: original})
                self.assertEqual(load_existing_results(self.path)[self.phone], original)

    def test_formula_like_customer_strings_are_preserved_as_literal_text(self):
        self.records[0].name = '=HYPERLINK("https://invalid.test","测试")'
        original = self.result(text="=1+1\n+重复原文\n@原始内容")
        original["error"] = "=SUM(1,2)"
        original["items"][0]["name"] = "=2+2"
        write_result(self.path, self.records, {self.phone: original})
        wb = load_workbook(self.path, data_only=False)
        self.addCleanup(wb.close)
        for ws in wb:
            for row in ws:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
        self.assertEqual(wb["输入数据"]["B2"].value, self.records[0].name)
        self.assertEqual(load_existing_results(self.path)[self.phone], original)

    def test_legacy_unsegmented_fulltext_and_english_status_are_read(self):
        original = self.result()
        legacy = {**original, "query_time": "2026-09-13T12:00:00"}
        write_result(self.path, self.records, {self.phone: legacy})
        wb = load_workbook(self.path)
        for sheet, cell in (("结果汇总", "D2"), ("营销详情", "D2"), ("运行日志", "A2")):
            self.assertEqual(wb[sheet][cell].value, original["query_time"])
        wb["结果汇总"]["C2"] = "success"
        wb["营销详情"].delete_cols(5)
        wb["推荐明细"].delete_cols(7)
        wb.save(self.path)
        wb.close()
        self.assertEqual(load_existing_results(self.path)[self.phone], original)

    def test_occupied_excel_preserves_previous_file_and_journal_keeps_new_result(self):
        previous = self.result(text="前一条已保存原文")
        write_result(self.path, self.records, {self.phone: previous})
        original_bytes = self.path.read_bytes()
        current = self.result(text="本次新采集内容")
        journal = append_result_journal(self.path, self.phone, current)
        with patch("excel_io.os.replace", side_effect=PermissionError("Excel 正在占用文件")):
            with self.assertRaises(PermissionError):
                write_result(self.path, self.records, {self.phone: current})
        self.assertEqual(self.path.read_bytes(), original_bytes)
        self.assertTrue(journal.exists())
        self.assertEqual(load_result_journal(self.path)[self.phone], current)
        self.assertEqual(list(self.path.parent.glob(".查询结果-*.xlsx")), [])

    def test_interrupted_workbook_save_preserves_old_file(self):
        write_result(self.path, self.records, {self.phone: self.result()})
        original_bytes = self.path.read_bytes()
        with patch("excel_io.Workbook.save", side_effect=OSError("磁盘不可写")):
            with self.assertRaises(OSError):
                write_result(self.path, self.records, {self.phone: self.result(text="新结果")})
        self.assertEqual(self.path.read_bytes(), original_bytes)
        self.assertEqual(list(self.path.parent.glob(".查询结果-*.xlsx")), [])

    def test_journal_latest_phone_result_wins_and_fsync_is_required(self):
        with patch("excel_io.os.fsync") as sync:
            journal = append_result_journal(self.path, self.phone, self.result(status="partial"))
            sync.assert_called_once()
        latest = self.result()
        self.assertEqual(append_result_journal(self.path, self.phone, latest), journal)
        self.assertEqual(journal, self.path.with_suffix(".jsonl"))
        self.assertEqual(load_result_journal(self.path), {self.phone: latest})
        with patch("excel_io.os.fsync", side_effect=OSError("同步失败")):
            with self.assertRaises(OSError):
                append_result_journal(self.path, self.phone, latest)

    def test_truncated_tail_recovers_complete_records_and_next_append_repairs_tail(self):
        complete = self.result()
        journal = append_result_journal(self.path, self.phone, complete)
        with journal.open("ab") as stream:
            stream.write(b'{"schema_version": 1, "phone": "138')
        self.assertEqual(load_result_journal(self.path), {self.phone: complete})
        next_phone = "13900000000"
        next_result = self.result(text="另一个测试号码全文")
        append_result_journal(self.path, next_phone, next_result)
        self.assertEqual(load_result_journal(self.path), {self.phone: complete, next_phone: next_result})

    def test_journal_interior_corruption_is_reported_without_overwriting(self):
        journal = append_result_journal(self.path, self.phone, self.result())
        with journal.open("ab") as stream:
            stream.write(b'{broken}\n')
        original_bytes = journal.read_bytes()
        with self.assertRaisesRegex(ValueError, "第 2 行损坏"):
            load_result_journal(self.path)
        self.assertEqual(journal.read_bytes(), original_bytes)


if __name__ == "__main__":
    unittest.main()
